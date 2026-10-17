"""Excel export of the dashboard for the selected period: one sheet per section."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.modules.demands.export import _safe

_FILL = PatternFill("solid", fgColor="E2E8F0")


def _sheet(wb: Workbook, title: str, header: list[tuple[str, int]], rows: list[list], first: bool = False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    for index, (name, width) in enumerate(header, start=1):
        cell = ws.cell(row=1, column=index, value=name)
        cell.font = Font(bold=True)
        cell.fill = _FILL
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A2"
    for row in rows:
        ws.append([_safe(v) for v in row])
    return ws


def build_dashboard_workbook(data: dict) -> bytes:
    wb = Workbook()
    p, t, pt = data["period"], data["totals"], data["previous_totals"]
    _sheet(wb, "Summary", [("Measure", 30), (p["label"], 20), (f"Previous ({data['previous']['label']})", 30)], [
        ["From", p["date_from"].strftime("%d %b %Y"), data["previous"]["date_from"].strftime("%d %b %Y")],
        ["To", p["date_to"].strftime("%d %b %Y"), data["previous"]["date_to"].strftime("%d %b %Y")],
        ["Orders (approved letters)", t["orders"], pt["orders"]],
        ["Vacancies", t["vacancies"], pt["vacancies"]],
        ["Clients (different employers)", t["clients"], pt["clients"]],
        ["New clients (first order ever)", t["new_clients"], pt["new_clients"]],
        ["Repeat clients", len(data["repeat_clients"]), None],
        [None, None, None],
        ["Pipeline right now", None, None],
        ["Drafts in progress", data["pipeline"]["drafts"], None],
        ["Waiting for admin approval", data["pipeline"]["waiting_approval"], None],
        ["Approved (all time)", data["pipeline"]["approved_all_time"], None],
    ], first=True)
    unit = "Day" if data["trend"]["unit"] == "day" else "Month"
    _sheet(wb, "Orders over time", [(unit, 14), ("Orders", 10), ("Vacancies", 12)],
           [[x["bucket"], x["orders"], x["vacancies"]] for x in data["trend"]["points"]])
    _sheet(wb, "Orders by client", [("Employer", 40), ("Orders", 10), ("Vacancies", 12)],
           [[c["employer"], c["orders"], c["vacancies"]] for c in data["by_client"]])
    _sheet(wb, "Repeat clients",
           [("Employer", 40), ("Months with 2+ orders", 22), ("Orders in those months", 22), ("Months", 40)],
           [[c["employer"], c["repeat_months"], c["orders_in_repeat_months"],
             ", ".join(f"{m['month']} ({m['orders']})" for m in c["months"])] for c in data["repeat_clients"]])
    _sheet(wb, "Orders by country", [("Country", 28), ("Orders", 10), ("Vacancies", 12)],
           [[c["country"], c["orders"], c["vacancies"]] for c in data["by_country"]])
    _sheet(wb, "Top positions", [("Position", 36), ("People needed", 14), ("Orders", 10)],
           [[x["position"], x["people"], x["orders"]] for x in data["positions"]])
    _sheet(wb, "Ad status", [("Ad status", 16), ("Orders", 10)],
           [[s["status"], s["orders"]] for s in data["ad_status"]])
    _sheet(wb, "Team", [("Created by", 28), ("Approved letters", 16), ("Vacancies", 12)],
           [[x["name"], x["orders"], x["vacancies"]] for x in data["team"]])
    out = BytesIO()
    wb.save(out)
    return out.getvalue()
