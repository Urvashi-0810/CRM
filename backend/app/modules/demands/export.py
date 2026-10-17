"""Excel export of the demand list (same filters as the screen)."""

from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.modules.demands.schemas import DemandSummary

IST = ZoneInfo("Asia/Kolkata")

# Same order as the screen (icons left out).
COLUMNS = [
    ("Approved date", 14),
    ("Country", 18),
    ("Demand ID", 12),
    ("Employer", 34),
    ("Vacancy", 10),
    ("Target timeline", 22),
    ("Assign to", 22),
    ("Status", 14),
    ("Created by", 22),
    ("Approved by", 22),
]
DRAFT_COLUMNS = [
    ("Created", 18),
    ("Demand ID", 12),
    ("Created by", 22),
    ("Country", 18),
    ("Employer", 34),
    ("Vacancy", 10),
]


def _ist(value):
    # Excel has no time zones: write India time without tz info.
    return value.astimezone(IST).replace(tzinfo=None) if value else None


def _safe(value):
    # Stop spreadsheet apps from treating user text as a formula (CSV/formula injection).
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def build_workbook(rows: list[DemandSummary], lifecycle: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Demand letters" if lifecycle == "approved" else "Drafts"
    columns = COLUMNS if lifecycle == "approved" else DRAFT_COLUMNS

    header_fill = PatternFill("solid", fgColor="E2E8F0")
    for index, (title, width) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=index, value=title)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A2"

    for r in rows:
        if lifecycle == "approved":
            values = [
                _ist(r.approved_at),
                r.country,
                r.demand_code,
                r.employer,
                r.vacancy,
                r.target_timeline,
                r.assign_to,
                r.progress_status,
                r.created_by.display_name,
                r.approved_by.display_name if r.approved_by else None,
            ]
        else:
            values = [
                _ist(r.created_at),
                r.demand_code,
                r.created_by.display_name,
                r.country,
                r.employer,
                r.vacancy,
            ]
        ws.append([_safe(v) for v in values])

    date_format = "dd-mmm-yyyy" if lifecycle == "approved" else "dd-mmm-yyyy hh:mm"
    for row in ws.iter_rows(min_row=2, min_col=1, max_col=1):
        for cell in row:
            cell.number_format = date_format
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
