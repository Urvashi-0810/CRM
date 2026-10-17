"""Turn the letter model into a .docx with the letterhead.

- Header: the letterhead header image across the full page width, then a line with
  "Demand Order ID : AA00001" on the left and "Date : <approval date>" on the right
  (blank until approval), plus a red DRAFT mark while the letter is not approved.
- Footer: the letterhead footer image across the full page width.
- Body: headings, paragraphs, bullet lists, tables and selected JD images.
- [[MISSING: X]] markers are shown highlighted in yellow so they are easy to find.
"""

import io
from dataclasses import dataclass
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

from app.modules.letters.letter_model import BOLD_RE, MISSING_RE, Letter

PAGE_WIDTH = Cm(21.0)
PAGE_HEIGHT = Cm(29.7)
SIDE_MARGIN = Cm(2.0)
BODY_WIDTH = PAGE_WIDTH - 2 * SIDE_MARGIN
BRAND_BLUE = RGBColor(0x1E, 0x3A, 0x8A)
DRAFT_RED = RGBColor(0xDC, 0x26, 0x26)


@dataclass
class LetterheadImages:
    header: bytes
    footer: bytes


@dataclass
class BuildInput:
    letter: Letter
    demand_code: str
    approval_date: date | None
    is_draft: bool
    letterhead: LetterheadImages | None
    images: dict[str, bytes]  # image_id -> image bytes


def _image_height(data: bytes, width_emu: int) -> int:
    with Image.open(io.BytesIO(data)) as img:
        w, h = img.size
    return int(width_emu * h / w)


def _full_width_picture(container, data: bytes) -> None:
    """A picture that spans the whole page width, ignoring the side margins."""
    p = container.paragraphs[0] if container.paragraphs and not container.paragraphs[0].text else container.add_paragraph()
    p.paragraph_format.left_indent = -SIDE_MARGIN
    p.paragraph_format.right_indent = -SIDE_MARGIN
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(io.BytesIO(data), width=PAGE_WIDTH)


def _no_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tbl_pr.append(borders)


def _add_rich_text(paragraph, text: str, *, bold: bool = False, size: Pt | None = None) -> None:
    """Add text with **bold** and [[MISSING: X]] support."""
    pos = 0
    tokens: list[tuple[str, str]] = []
    pattern = f"{BOLD_RE.pattern}|{MISSING_RE.pattern}"
    import re

    for m in re.finditer(pattern, text):
        if m.start() > pos:
            tokens.append(("text", text[pos:m.start()]))
        if m.group(1) is not None:
            tokens.append(("bold", m.group(1)))
        else:
            tokens.append(("missing", m.group(2)))
        pos = m.end()
    if pos < len(text):
        tokens.append(("text", text[pos:]))

    for kind, value in tokens:
        if kind == "missing":
            run = paragraph.add_run(f"[{value} — to be filled]")
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            run.bold = True
        else:
            run = paragraph.add_run(value)
            run.bold = bold or kind == "bold"
        if size:
            run.font.size = size


def _shade(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _header(section, data: BuildInput) -> None:
    header = section.header
    header.is_linked_to_previous = False
    if data.letterhead:
        _full_width_picture(header, data.letterhead.header)
    info = header.add_table(rows=1, cols=2, width=BODY_WIDTH)
    _no_borders(info)
    left, right = info.rows[0].cells
    lp = left.paragraphs[0]
    _add_rich_text(lp, "Demand Order ID : ", bold=True, size=Pt(10))
    run = lp.add_run(data.demand_code)
    run.bold, run.font.size, run.font.color.rgb = True, Pt(10), BRAND_BLUE
    rp = right.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _add_rich_text(rp, "Date : ", bold=True, size=Pt(10))
    if data.approval_date:
        run = rp.add_run(data.approval_date.strftime("%d %B %Y"))
        run.bold, run.font.size, run.font.color.rgb = True, Pt(10), BRAND_BLUE
    if data.is_draft:
        dp = header.add_paragraph()
        dp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = dp.add_run("DRAFT — NOT APPROVED")
        run.bold, run.font.size, run.font.color.rgb = True, Pt(10), DRAFT_RED


def _footer(section, data: BuildInput) -> None:
    footer = section.footer
    footer.is_linked_to_previous = False
    if data.letterhead:
        _full_width_picture(footer, data.letterhead.footer)


def build_docx(data: BuildInput) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width, section.page_height = PAGE_WIDTH, PAGE_HEIGHT
    section.left_margin = section.right_margin = SIDE_MARGIN
    if data.letterhead:
        header_h = _image_height(data.letterhead.header, PAGE_WIDTH)
        footer_h = _image_height(data.letterhead.footer, PAGE_WIDTH)
        section.header_distance = Cm(0)
        section.footer_distance = Cm(0)
        section.top_margin = header_h + Cm(1.6)
        section.bottom_margin = footer_h + Cm(0.6)
    else:
        section.top_margin = Cm(3.0)
        section.bottom_margin = Cm(2.0)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)

    _header(section, data)
    _footer(section, data)

    letter = data.letter
    if letter.title:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_rich_text(p, letter.title, bold=True, size=Pt(14))
        for r in p.runs:
            if r.font.highlight_color is None:
                r.font.color.rgb = BRAND_BLUE

    for block in letter.blocks:
        if block.type == "heading":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.keep_with_next = True
            _add_rich_text(p, block.text or "", bold=True, size=Pt({1: 13, 2: 11.5, 3: 10.5}[block.level or 2]))
            for r in p.runs:
                if r.font.highlight_color is None:
                    r.font.color.rgb = BRAND_BLUE
        elif block.type == "paragraph":
            for line in (block.text or "").split("\n"):
                _add_rich_text(doc.add_paragraph(), line)
        elif block.type == "bullets":
            for item in block.items or []:
                _add_rich_text(doc.add_paragraph(style="List Bullet"), item)
        elif block.type == "table":
            columns, rows = block.columns or [], block.rows or []
            has_header = any(c.strip() for c in columns)
            table = doc.add_table(rows=(1 if has_header else 0) + len(rows), cols=len(columns) or len(rows[0]))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            offset = 0
            if has_header:
                for i, col in enumerate(columns):
                    cell = table.rows[0].cells[i]
                    _add_rich_text(cell.paragraphs[0], col, bold=True, size=Pt(10))
                    _shade(cell, "DBEAFE")
                offset = 1
            for r_i, row in enumerate(rows):
                for c_i, value in enumerate(row):
                    _add_rich_text(table.rows[r_i + offset].cells[c_i].paragraphs[0], value, size=Pt(10))
            doc.add_paragraph()
        elif block.type == "image":
            image = data.images.get(block.image_id or "")
            if image is None:
                continue
            with Image.open(io.BytesIO(image)) as img:
                w_px = img.size[0]
            # ~96 dpi on screen; never wider than the text column
            width = min(BODY_WIDTH, Cm(w_px / 96 * 2.54))
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(io.BytesIO(image), width=width)
            if block.caption:
                cp = doc.add_paragraph()
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _add_rich_text(cp, block.caption, size=Pt(9))

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
