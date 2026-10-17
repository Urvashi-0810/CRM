"""Working with existing .docx files: update the header stamp, read a letter back.

- stamp_docx: change only the header's Date and DRAFT mark, leaving everything else in the
  file untouched (used on approval and when restoring an old version).
- docx_to_letter: read an edited letter (downloaded, changed in Word, uploaded back) into the
  letter model, so the preview shows it and the AI's next edit starts from it.
"""

import io
import re
from dataclasses import dataclass
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from PIL import Image

from app.modules.letters.letter_model import Block, Letter

DRAFT_TEXT = "DRAFT — NOT APPROVED"
CODE_RE = re.compile(r"\b[A-Z]{2}\d{5}\b")
FILL_RE = re.compile(r"\[([^\[\]]+?) — to be filled\]")
BRAND_BLUE = RGBColor(0x1E, 0x3A, 0x8A)
DRAFT_RED = RGBColor(0xDC, 0x26, 0x26)


def _header_parts(doc):
    for section in doc.sections:
        for header in (section.header, section.first_page_header, section.even_page_header):
            if header is not None and not header.is_linked_to_previous:
                yield header


def header_codes(data: bytes) -> set[str]:
    """Demand codes written in the file's header(s)."""
    doc = Document(io.BytesIO(data))
    found: set[str] = set()
    for header in _header_parts(doc):
        texts = [p.text for p in header.paragraphs]
        for table in header.tables:
            texts += [c.text for c in table._cells]
        for t in texts:
            found.update(CODE_RE.findall(t))
    return found


def stamp_docx(data: bytes, *, approval_date: date | None, is_draft: bool) -> bytes | None:
    """Set the header's Date and DRAFT mark in place. Returns None if the header layout made
    by the app (a "Date :" line) can't be found, e.g. the user removed the header."""
    doc = Document(io.BytesIO(data))
    stamped = False
    for header in _header_parts(doc):
        date_paragraphs = [
            p for table in header.tables for cell in table._cells for p in cell.paragraphs
            if p.text.strip().startswith("Date :")
        ] + [p for p in header.paragraphs if p.text.strip().startswith("Date :")]
        for p in date_paragraphs:
            stamped = True
            runs = p.runs
            label = next((r for r in runs if "Date :" in r.text), None)
            for r in runs:
                if r is not label:
                    r._element.getparent().remove(r._element)
            if label is not None:
                label.text = "Date : "
            if approval_date:
                run = p.add_run(approval_date.strftime("%d %B %Y"))
                run.bold, run.font.size, run.font.color.rgb = True, Pt(10), BRAND_BLUE

        drafts = [p for p in header.paragraphs if DRAFT_TEXT in p.text]
        if not is_draft:
            for p in drafts:
                p._element.getparent().remove(p._element)
        elif date_paragraphs and not drafts:
            dp = header.add_paragraph()
            dp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run = dp.add_run(DRAFT_TEXT)
            run.bold, run.font.size, run.font.color.rgb = True, Pt(10), DRAFT_RED
    if not stamped:
        return None
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


@dataclass
class FoundImage:
    placeholder: str  # "upload-1", replaced by the stored image id
    data: bytes
    mime: str
    width: int
    height: int


def _paragraph_text(p) -> str:
    """Text with **bold** and [[MISSING: X]] markers restored."""
    pieces: list[str] = []
    for run in p.runs:
        text = run.text
        if not text:
            continue
        if run.font.highlight_color is not None and FILL_RE.search(text):
            text = FILL_RE.sub(lambda m: f"[[MISSING: {m.group(1)}]]", text)
            pieces.append(text)
        elif run.bold and text.strip():
            lead, core, trail = re.match(r"(\s*)(.*?)(\s*)$", text, re.S).groups()
            pieces.append(f"{lead}**{core}**{trail}")
        else:
            pieces.append(text)
    joined = "".join(pieces).replace("****", "")
    return FILL_RE.sub(lambda m: f"[[MISSING: {m.group(1)}]]", joined)


def _all_bold(p) -> bool:
    runs = [r for r in p.runs if r.text.strip()]
    return bool(runs) and all(r.bold for r in runs)


def docx_to_letter(data: bytes) -> tuple[Letter, list[FoundImage]]:
    doc = Document(io.BytesIO(data))
    blocks: list[Block] = []
    images: list[FoundImage] = []
    title: str | None = None
    rels = doc.part.rels

    def flush_bullets(items: list[str]) -> None:
        if items:
            blocks.append(Block(type="bullets", items=list(items)))
            items.clear()

    bullets: list[str] = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            from docx.text.paragraph import Paragraph

            p = Paragraph(child, doc)
            for blip in child.iter(qn("a:blip")):
                rid = blip.get(qn("r:embed"))
                part = rels[rid].target_part if rid in rels else None
                if part is None:
                    continue
                try:
                    with Image.open(io.BytesIO(part.blob)) as img:
                        fmt, (w, h) = img.format, img.size
                except Exception:
                    continue
                mime = {"PNG": "image/png", "JPEG": "image/jpeg", "GIF": "image/gif", "WEBP": "image/webp"}.get(fmt)
                if mime is None:
                    continue
                flush_bullets(bullets)
                placeholder = f"upload-{len(images) + 1}"
                images.append(FoundImage(placeholder, part.blob, mime, w, h))
                blocks.append(Block(type="image", image_id=placeholder))
            text = _paragraph_text(p).strip()
            if not text:
                continue
            style = (p.style.name or "").lower() if p.style is not None else ""
            if style.startswith("list"):
                bullets.append(text)
                continue
            flush_bullets(bullets)
            plain = text.replace("**", "")
            if style.startswith("title") or (
                title is None and not blocks and _all_bold(p)
                and p.alignment == WD_ALIGN_PARAGRAPH.CENTER
            ):
                title = plain
            elif style.startswith("heading") or (_all_bold(p) and len(plain) <= 120):
                level = 1 if style.startswith("heading 1") else 2
                blocks.append(Block(type="heading", text=plain, level=level))
            elif blocks and blocks[-1].type == "image" and p.alignment == WD_ALIGN_PARAGRAPH.CENTER and len(plain) <= 200:
                blocks[-1] = blocks[-1].model_copy(update={"caption": plain})
            else:
                blocks.append(Block(type="paragraph", text=text))
        elif child.tag == qn("w:tbl"):
            flush_bullets(bullets)
            from docx.table import Table

            table = Table(child, doc)
            rows: list[list[str]] = []
            for row in table.rows:
                cells, previous = [], None
                for cell in row.cells:
                    if previous is not None and cell._tc is previous:
                        continue  # merged cell repeated by python-docx
                    previous = cell._tc
                    cells.append(" ".join(_paragraph_text(cp).strip() for cp in cell.paragraphs).strip())
                rows.append(cells)
            if rows:
                blocks.append(Block(type="table", columns=[c.replace("**", "") for c in rows[0]], rows=rows[1:]))
    flush_bullets(bullets)
    return Letter(title=title, blocks=blocks), images


def replace_image_ids(letter: Letter, mapping: dict[str, str]) -> Letter:
    blocks = [
        b.model_copy(update={"image_id": mapping.get(b.image_id, b.image_id)}) if b.type == "image" else b
        for b in letter.blocks
    ]
    return letter.model_copy(update={"blocks": blocks})

