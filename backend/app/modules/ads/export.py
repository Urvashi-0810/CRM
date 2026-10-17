"""Excel export of the Create Advertisement table (same filters as the screen)."""

from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.modules.demands.export import _safe
from app.modules.demands.schemas import DemandSummary

IST = ZoneInfo("Asia/Kolkata")

# Same order as the screen (icons left out).
COLUMNS = [
    ("Approved on", 14),
    ("Country", 18),
    ("Demand ID", 12),
    ("Employer", 34),
    ("Vacancy", 10),
    ("Target timeline", 22),
    ("Ad assigned to", 22),
    ("Ad status", 14),
    ("Ad status changed by", 22),
    ("Ad status changed on", 18),
]


def _ist(value):
    return value.astimezone(IST).replace(tzinfo=None) if value else None


def build_ads_workbook(rows: list[DemandSummary]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Advertisements"
    fill = PatternFill("solid", fgColor="E2E8F0")
    for index, (title, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=index, value=title)
        cell.font = Font(bold=True)
        cell.fill = fill
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A2"
    for r in rows:
        ws.append([_safe(v) for v in (
            _ist(r.approved_at), r.country, r.demand_code, r.employer, r.vacancy, r.target_timeline,
            r.ad_assigned_to, r.ad_status, r.ad_updated_by.display_name if r.ad_updated_by else None,
            _ist(r.ad_updated_at),
        )])
    for row in ws.iter_rows(min_row=2):
        row[0].number_format = "dd-mmm-yyyy"
        row[9].number_format = "dd-mmm-yyyy hh:mm"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(rows) + 1}"
    out = BytesIO()
    wb.save(out)
    return out.getvalue()
