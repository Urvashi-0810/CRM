"""Chat → worker → AI → letter, with a scripted fake AI (no network)."""

import io
import json
from pathlib import Path

import pytest
from docx import Document
from PIL import Image
from sqlalchemy import select, text

from app.db import get_sessionmaker
from app.modules.ai_admin.models import LlmCall
from app.modules.jobs.queue import claim_next
from app.modules.jobs.runner import run_job
from app.modules.llm import gateway
from app.modules.llm.base import LlmError, LlmResult
from app.modules.storage.local import LocalStorage
from tests.conftest import login, unlock_admin

FIX = Path(__file__).parent / "fixtures"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def blank_block(**kw):
    return {"text": None, "level": None, "items": None, "columns": None, "rows": None,
            "image_id": None, "caption": None, **kw}


def reply(action, message="", **kw):
    body = {"action": action, "message": message, "questions": [], "missing_fields": [],
            "letter": None, "extracted": None}
    body.update(kw)
    return json.dumps(body)


ASK = reply("ask", "Before I write the letter:", questions=["Which country is the job in?"])
LETTER = reply(
    "letter", "Letter ready. Contract period is missing.",
    missing_fields=["Contract period"],
    letter={"title": "Demand Letter", "blocks": [
        blank_block(type="paragraph", text="**Employer:** Mosanada Facility Management Services"),
        blank_block(type="table", columns=["Project", "Position", "Qty"],
                    rows=[["Aspire Zone", "HVAC/Mechanical Engineer", "2"], ["MSY", "Housekeeping Officer", "10"]]),
        blank_block(type="paragraph", text="Contract: [[MISSING: Contract period]]"),
    ]},
    extracted={"country": "Qatar", "employer": "Mosanada Facility Management Services", "vacancy": 12,
               "trades": [{"project": "Aspire Zone", "trade": "HVAC/Mechanical Engineer", "quantity": 2},
                          {"project": "MSY", "trade": "Housekeeping Officer", "quantity": 10}]},
)


class ScriptedAdapter:
    provider = "claude"

    def __init__(self, replies):
        self.replies = list(replies)
        self.requests = []

    async def generate(self, request):
        self.requests.append(request)
        item = self.replies.pop(0)
        if isinstance(item, Exception):
            raise item
        return LlmResult(text=item, input_tokens=1000, output_tokens=200, cache_read_tokens=500)


@pytest.fixture
def ai(monkeypatch):
    holder = {}

    def install(*replies):
        adapter = ScriptedAdapter(replies)
        holder["adapter"] = adapter
        monkeypatch.setitem(gateway.ADAPTER_FACTORIES, "claude", lambda key, model: adapter)
        return adapter

    async def no_sleep(_):
        return None

    monkeypatch.setattr(gateway.asyncio, "sleep", no_sleep)
    return install


async def configure_ai(client, letterhead=True):
    assert (await unlock_admin(client)).status_code == 204
    created = await client.post("/api/v1/admin/llm-providers",
                                json={"provider": "claude", "model": "claude-sonnet-5-5", "api_key": "sk-ant-test-1234"})
    assert created.status_code == 201, created.text
    assert created.json()["key_last4"] == "1234" and "api_key" not in created.json()
    assert (await client.post(f"/api/v1/admin/llm-providers/{created.json()['id']}/activate")).status_code == 200
    if letterhead:
        up = await client.post("/api/v1/admin/letterheads", files={
            "header": ("header.png", (FIX / "test_header.png").read_bytes(), "image/png"),
            "footer": ("footer.png", (FIX / "test_footer.png").read_bytes(), "image/png"),
        })
        assert up.status_code == 201, up.text


async def setup(client, make_user, admin_password_set):
    await make_user("priya", "Priya Shah")
    await login(client, "priya")
    await configure_ai(client)
    demand = (await client.post("/api/v1/demands")).json()
    return demand


async def send(client, demand_id, text="", files=()):
    multipart = [("files", (name, data, mime)) for name, data, mime, _ in files]
    data = {"text": text, "kinds": [kind for *_, kind in files]}
    return await client.post(f"/api/v1/demands/{demand_id}/messages", data=data, files=multipart or None)


async def work(settings):
    """Run the next job the way the worker does."""
    async with get_sessionmaker()() as db:
        job_id = await claim_next(db, 300)
    assert job_id is not None
    await run_job(job_id, sessionmaker=get_sessionmaker(), storage=LocalStorage(settings.storage_root),
                  settings=settings)
    return job_id


def inquiry_file(name="inquiry1_mosanada.docx"):
    return (name, (FIX / name).read_bytes(), DOCX, "inquiry")


async def test_ask_then_letter_then_download(client, make_user, admin_password_set, ai, settings):
    adapter = ai(ASK, LETTER)
    demand = await setup(client, make_user, admin_password_set)

    sent = await send(client, demand["id"], "Please prepare the demand letter",
                      [inquiry_file(), ("jd.xlsx", (FIX / "jd_list_of_resources.xlsx").read_bytes(),
                                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "jd")])
    assert sent.status_code == 202, sent.text
    job = sent.json()["job"]
    assert job["state"] == "queued"
    assert [f["kind"] for f in sent.json()["message"]["files"]] == ["inquiry", "jd"]

    # A second message while the first is being worked on is refused.
    busy = await send(client, demand["id"], "hello?")
    assert busy.status_code == 409 and busy.json()["type"] == "urn:crm:problem:job-in-progress"

    await work(settings)
    assert (await client.get(f"/api/v1/jobs/{job['id']}")).json()["state"] == "succeeded"
    messages = (await client.get(f"/api/v1/demands/{demand['id']}/messages")).json()["items"]
    assert messages[-1]["kind"] == "questions"
    assert messages[-1]["data"]["questions"] == ["Which country is the job in?"]
    assert (await client.get(f"/api/v1/demands/{demand['id']}/letter")).status_code == 404

    # What the AI received: rules, client prompt, both attachments as text, the user's words.
    first = adapter.requests[0]
    assert "NEVER INVENT" in first.system and "CLIENT INSTRUCTIONS" in first.system
    sent_text = "\n".join(p.text or "" for p in first.turns[0].parts)
    assert "Demand inquiry: inquiry1_mosanada.docx" in sent_text and "Genset Technician" in sent_text
    assert sent_text.endswith("Please prepare the demand letter")

    await send(client, demand["id"], "The job is in Qatar.")
    await work(settings)

    letter = (await client.get(f"/api/v1/demands/{demand['id']}/letter")).json()
    assert letter["version_no"] == 1 and letter["source"] == "generated"
    assert letter["missing_fields"] == ["Contract period"]
    assert letter["letterhead"]["header_file_id"]
    last = (await client.get(f"/api/v1/demands/{demand['id']}/messages")).json()["items"][-1]
    assert last["kind"] == "letter" and last["data"]["version_no"] == 1

    # The AI's extraction filled the list fields and trades.
    detail = (await client.get(f"/api/v1/demands/{demand['id']}")).json()
    assert (detail["country"], detail["employer"], detail["vacancy"]) == ("Qatar", "Mosanada Facility Management Services", 12)
    assert [t["trade"] for t in detail["trades"]] == ["HVAC/Mechanical Engineer", "Housekeeping Officer"]
    assert detail["latest_version_no"] == 1

    download = await client.get(f"/api/v1/demands/{demand['id']}/docx")
    assert download.status_code == 200 and download.headers["content-type"] == DOCX
    assert f"{demand['demand_code']}_v1.docx" in download.headers["content-disposition"]
    doc = Document(io.BytesIO(download.content))
    assert "Housekeeping Officer" in [c.text for c in doc.tables[0]._cells]

    # Every AI call was logged with tokens and cost.
    async with get_sessionmaker()() as db:
        calls = list((await db.execute(select(LlmCall))).scalars())
    assert len(calls) == 2 and all(c.outcome == "ok" and c.cost_usd > 0 for c in calls)


async def test_pasted_inquiry_saved_as_docx(client, make_user, admin_password_set, ai, settings):
    adapter = ai(ASK)
    demand = await setup(client, make_user, admin_password_set)
    mail = "Dear Sir,\r\nWe need 12 Genset Technicians for Qatar.\x0bSalary: QAR 2500\x00\n\nRegards"

    sent = await client.post(f"/api/v1/demands/{demand['id']}/messages",
                             data={"text": "", "inquiry_text": mail})
    assert sent.status_code == 202, sent.text
    [f] = sent.json()["message"]["files"]
    assert (f["kind"], f["name"]) == ("inquiry", "Demand inquiry (pasted mail).docx")

    download = await client.get(f"/api/v1/files/{f['id']}")
    assert download.status_code == 200
    lines = [p.text for p in Document(io.BytesIO(download.content)).paragraphs]
    assert lines == ["Dear Sir,", "We need 12 Genset Technicians for Qatar.", "Salary: QAR 2500", "", "Regards"]

    await work(settings)
    sent_text = "\n".join(p.text or "" for p in adapter.requests[0].turns[0].parts)
    assert "Demand inquiry: Demand inquiry (pasted mail).docx" in sent_text
    assert "Genset Technicians for Qatar" in sent_text

    too_long = await client.post(f"/api/v1/demands/{demand['id']}/messages",
                                 data={"inquiry_text": "x" * 100_001})
    assert too_long.status_code == 422 and too_long.json()["type"] == "urn:crm:problem:inquiry-too-long"


async def test_prompt_edit_sends_current_letter_and_keeps_manual_fields(
    client, make_user, admin_password_set, ai, settings
):
    edited = json.loads(LETTER)
    edited["letter"]["blocks"][2]["text"] = "Contract: 2 years"
    edited["missing_fields"] = []
    edited["extracted"]["employer"] = "MFMS (AI)"
    adapter = ai(LETTER, json.dumps(edited))
    demand = await setup(client, make_user, admin_password_set)
    await send(client, demand["id"], "", [inquiry_file()])
    await work(settings)

    # A person corrects the employer by hand: the AI must not overwrite it afterwards.
    current = await client.get(f"/api/v1/demands/{demand['id']}")
    await client.patch(f"/api/v1/demands/{demand['id']}", json={"employer": "Mosanada FMS"},
                       headers={"If-Match": current.headers["etag"]})

    await send(client, demand["id"], "Contract period is 2 years")
    await work(settings)

    context = adapter.requests[1].turns[-1].parts[0].text
    assert context.startswith("CURRENT LETTER (version 1)") and "[[MISSING: Contract period]]" in context
    letter = (await client.get(f"/api/v1/demands/{demand['id']}/letter")).json()
    assert letter["version_no"] == 2 and letter["source"] == "prompt_edit" and letter["missing_fields"] == []
    assert (await client.get(f"/api/v1/demands/{demand['id']}")).json()["employer"] == "Mosanada FMS"


async def test_off_topic_gets_fixed_refusal(client, make_user, admin_password_set, ai, settings):
    ai(reply("refuse", "whatever the model says"))
    demand = await setup(client, make_user, admin_password_set)
    await send(client, demand["id"], "What's the weather in Goa?")
    await work(settings)
    last = (await client.get(f"/api/v1/demands/{demand['id']}/messages")).json()["items"][-1]
    assert last["kind"] == "refusal"
    assert last["body"].startswith("I can only help with demand letters")


async def test_rate_limit_retried_then_succeeds_and_each_attempt_logged(
    client, make_user, admin_password_set, ai, settings
):
    ai(LlmError("rate_limited", "slow down", retry_after=1), LlmError("overloaded", "busy"), ASK)
    demand = await setup(client, make_user, admin_password_set)
    await send(client, demand["id"], "", [inquiry_file()])
    await work(settings)
    async with get_sessionmaker()() as db:
        outcomes = [c.outcome for c in (await db.execute(select(LlmCall).order_by(LlmCall.id))).scalars()]
    assert outcomes == ["rate_limited", "overloaded", "ok"]


async def test_bad_api_key_fails_job_with_clear_message(client, make_user, admin_password_set, ai, settings):
    ai(LlmError("auth", "The Claude API key was rejected. Check it on the admin page."))
    demand = await setup(client, make_user, admin_password_set)
    sent = await send(client, demand["id"], "", [inquiry_file()])
    await work(settings)
    job = (await client.get(f"/api/v1/jobs/{sent.json()['job']['id']}")).json()
    assert job["state"] == "failed" and "API key was rejected" in job["error"]
    last = (await client.get(f"/api/v1/demands/{demand['id']}/messages")).json()["items"][-1]
    assert last["kind"] == "error"
    # The chat is usable again afterwards.
    ai(ASK)
    assert (await send(client, demand["id"], "retry")).status_code == 202


async def test_invalid_ai_output_retried_once(client, make_user, admin_password_set, ai, settings):
    ai("not json at all", ASK)
    demand = await setup(client, make_user, admin_password_set)
    await send(client, demand["id"], "", [inquiry_file()])
    await work(settings)
    last = (await client.get(f"/api/v1/demands/{demand['id']}/messages")).json()["items"][-1]
    assert last["kind"] == "questions"


async def test_jd_pictures_offered_to_ai_and_embedded(client, make_user, admin_password_set, ai, settings):
    doc = Document()
    doc.add_paragraph("Uniform reference")
    pic = io.BytesIO()
    Image.new("RGB", (400, 300), (200, 30, 30)).save(pic, "PNG")
    doc.add_picture(io.BytesIO(pic.getvalue()))
    jd = io.BytesIO()
    doc.save(jd)

    # First reply just asks; the picture's id (img-<file id>) is read from what the AI received.
    adapter = ai(ASK)
    demand = await setup(client, make_user, admin_password_set)
    await send(client, demand["id"], "", [inquiry_file(), ("jd.docx", jd.getvalue(), DOCX, "jd")])
    await work(settings)
    texts = "\n".join(p.text or "" for p in adapter.requests[0].turns[0].parts)
    image_id = texts.split("image_id=")[1].split()[0]
    assert any(p.kind == "image" for p in adapter.requests[0].turns[0].parts)

    with_image = json.loads(LETTER)
    with_image["letter"]["blocks"].append(blank_block(type="image", image_id=image_id, caption="Uniform"))
    with_image["letter"]["blocks"].append(blank_block(type="image", image_id="img-999999", caption="made up"))
    ai(json.dumps(with_image))
    await send(client, demand["id"], "Go ahead")
    await work(settings)
    letter = (await client.get(f"/api/v1/demands/{demand['id']}/letter")).json()
    images = [b["image_id"] for b in letter["content"]["blocks"] if b["type"] == "image"]
    assert images == [image_id]  # the invented id was dropped
    file_id = int(image_id.split("-")[1])
    served = await client.get(f"/api/v1/files/{file_id}")
    assert served.status_code == 200 and served.headers["content-type"] == "image/png"


# ---------- upload rules ----------


async def test_file_too_large_message(client, make_user, admin_password_set, ai, settings):
    ai()
    demand = await setup(client, make_user, admin_password_set)
    big = b"0" * (settings.max_file_mb * 1_000_000 + 1)
    response = await send(client, demand["id"], "", [("big.txt", big, "text/plain", "jd")])
    assert response.status_code == 413
    body = response.json()
    assert body["title"] == "File too large" and '"big.txt" is larger than 10.0 MB' in body["detail"]


async def test_upload_rules(client, make_user, admin_password_set, ai):
    ai()
    demand = await setup(client, make_user, admin_password_set)
    no_label = await client.post(f"/api/v1/demands/{demand['id']}/messages",
                                 files=[("files", ("a.docx", b"x", DOCX))])
    assert no_label.status_code == 422 and no_label.json()["type"] == "urn:crm:problem:file-label-missing"
    wrong_type = await send(client, demand["id"], "", [("inq.xlsx", b"x", "application/octet-stream", "inquiry")])
    assert wrong_type.status_code == 415
    old = await send(client, demand["id"], "", [("old.doc", b"\xd0\xcf", "application/msword", "jd")])
    assert old.status_code == 415 and "save it as .docx" in old.json()["detail"]
    empty = await send(client, demand["id"], "")
    assert empty.status_code == 422


async def test_chat_needs_ai_configured(client, make_user):
    await make_user("priya")
    await login(client, "priya")
    demand = (await client.post("/api/v1/demands")).json()
    response = await send(client, demand["id"], "hello")
    assert response.status_code == 409 and response.json()["type"] == "urn:crm:problem:ai-not-configured"
    status = (await client.get("/api/v1/ai/status")).json()
    assert status["configured"] is False


async def test_delete_draft_with_letter_and_blocked_while_working(client, make_user, admin_password_set, ai, settings):
    ai(LETTER)
    demand = await setup(client, make_user, admin_password_set)
    await send(client, demand["id"], "", [inquiry_file()])
    blocked = await client.delete(f"/api/v1/demands/{demand['id']}")
    assert blocked.status_code == 409
    await work(settings)
    # Usage records survive the delete (money already spent).
    assert (await client.delete(f"/api/v1/demands/{demand['id']}")).status_code == 204
    async with get_sessionmaker()() as db:
        count = (await db.execute(text("select count(*) from llm_calls where demand_id is null"))).scalar_one()
    assert count == 1


async def test_crashed_worker_job_is_picked_up_again(client, make_user, admin_password_set, ai, settings):
    ai(ASK)
    demand = await setup(client, make_user, admin_password_set)
    await send(client, demand["id"], "", [inquiry_file()])
    async with get_sessionmaker()() as db:
        first = await claim_next(db, 300)
        # Simulate the worker dying: lease expires.
        await db.execute(text("update generation_jobs set lease_until = now() - interval '1 second'"))
        await db.commit()
        again = await claim_next(db, 300)
    assert again == first
    await run_job(again, sessionmaker=get_sessionmaker(), storage=LocalStorage(settings.storage_root),
                  settings=settings)
    last = (await client.get(f"/api/v1/demands/{demand['id']}/messages")).json()["items"][-1]
    assert last["kind"] == "questions"


# ---------- admin AI settings ----------


async def test_admin_ai_settings(client, make_user, admin_password_set):
    await make_user("ravi")
    await login(client, "ravi")
    assert (await client.get("/api/v1/admin/llm-providers")).status_code == 403
    await unlock_admin(client)

    listing = (await client.get("/api/v1/admin/llm-providers")).json()
    assert listing["providers"] == ["claude", "gemini"]  # test provider hidden unless enabled
    no_key = await client.post("/api/v1/admin/llm-providers", json={"provider": "gemini", "model": "gemini-3.8-flash"})
    assert no_key.status_code == 422
    test_provider = await client.post("/api/v1/admin/llm-providers", json={"provider": "test", "model": "test"})
    assert test_provider.status_code == 422

    prompts = (await client.get("/api/v1/admin/prompts")).json()["items"]
    assert prompts[0]["is_active"]  # default prompt from the migration
    saved = await client.post("/api/v1/admin/prompts", json={"body": "Client prompt: always include a positions table."})
    assert saved.status_code == 201 and saved.json()["is_active"]
    items = (await client.get("/api/v1/admin/prompts")).json()["items"]
    assert [p["is_active"] for p in items[:2]] == [True, False]
    back = await client.post(f"/api/v1/admin/prompts/{items[1]['id']}/activate")
    assert back.status_code == 200

    bad = await client.post("/api/v1/admin/letterheads", files={
        "header": ("h.gif", b"GIF89a", "image/gif"), "footer": ("f.png", b"x", "image/png")})
    assert bad.status_code in (415, 422)
