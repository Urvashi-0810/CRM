"""Create Advertisement: approved letters only, ad status/assignee, exports for Canva and Figma."""

import io
import zipfile

from openpyxl import load_workbook
from PIL import Image

from app.modules.ads import router as ads_router
from tests.test_chat_flow import LETTER, ai, send, work  # noqa: F401  (ai is a fixture)
from tests.test_milestone4 import letter_ready, submit_and_approve


def two_page_pdf() -> bytes:
    pages = [Image.new("RGB", (595, 842), c) for c in ("white", "lightblue")]
    out = io.BytesIO()
    pages[0].save(out, "PDF", save_all=True, append_images=pages[1:])
    return out.getvalue()


async def approved_letter(client, make_user, admin_password_set, ai, settings, replies=(LETTER,)):
    demand = await letter_ready(client, make_user, admin_password_set, ai, settings, replies=replies)
    assert (await submit_and_approve(client, demand["id"], 1)).status_code == 200
    return demand


async def test_only_approved_letters_and_ad_fields(client, make_user, admin_password_set, ai, settings):
    demand = await approved_letter(client, make_user, admin_password_set, ai, settings)
    draft = (await client.post("/api/v1/demands")).json()

    page = (await client.get("/api/v1/ads")).json()
    assert [i["demand_code"] for i in page["items"]] == [demand["demand_code"]]
    row = page["items"][0]
    assert row["ad_status"] == "Not started" and row["ad_assigned_to"] is None and row["ad_updated_by"] is None

    etag_before = (await client.get(f"/api/v1/demands/{demand['id']}")).json()["etag"]
    changed = await client.patch(f"/api/v1/ads/{demand['id']}", json={"ad_status": "In progress", "ad_assigned_to": " Ravi "})
    assert changed.status_code == 200, changed.text
    body = changed.json()
    assert (body["ad_status"], body["ad_assigned_to"]) == ("In progress", "Ravi")
    assert body["ad_updated_by"]["display_name"] == "Priya Shah" and body["ad_updated_at"]

    # Filters and options follow the ad fields.
    assert (await client.get("/api/v1/ads?ad_status=Done")).json()["total"] == 0
    assert (await client.get("/api/v1/ads?ad_status=In progress&ad_assigned_to=Ravi")).json()["total"] == 1
    assert (await client.get("/api/v1/ads/filter-options")).json()["ad_assignees"] == ["Ravi"]

    # The letter itself is untouched (same ETag as before the ad changes), drafts can't be reached, unknown statuses are refused.
    assert (await client.get(f"/api/v1/demands/{demand['id']}")).json()["etag"] == etag_before
    assert (await client.patch(f"/api/v1/ads/{draft['id']}", json={"ad_status": "Done"})).status_code == 404
    assert (await client.patch(f"/api/v1/ads/{demand['id']}", json={"ad_status": "Finished"})).status_code == 422
    cleared = await client.patch(f"/api/v1/ads/{demand['id']}", json={"ad_assigned_to": ""})
    assert cleared.json()["ad_assigned_to"] is None and cleared.json()["ad_status"] == "In progress"

    sheet = load_workbook(io.BytesIO((await client.get("/api/v1/ads/export")).content)).active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0][:3] == ("Approved on", "Country", "Demand ID") and rows[0][7] == "Ad status"
    assert len(rows) == 2 and rows[1][7] == "In progress"


async def test_exports_for_canva_and_figma(client, make_user, admin_password_set, ai, settings, monkeypatch):
    demand = await approved_letter(client, make_user, admin_password_set, ai, settings, replies=(LETTER, LETTER))
    url = f"/api/v1/ads/{demand['id']}/export"
    calls = []

    async def fake_pdf(docx, soffice):
        calls.append(len(docx))
        return two_page_pdf()

    monkeypatch.setattr(ads_router, "docx_to_pdf", fake_pdf)

    docx = await client.get(f"{url}?tool=canva&format=docx")
    assert docx.status_code == 200 and docx.content[:2] == b"PK"
    assert f"{demand['demand_code']}_v2.docx" in docx.headers["content-disposition"]  # v2 = approval stamp

    wrong = await client.get(f"{url}?tool=figma&format=docx")
    assert wrong.status_code == 422 and wrong.json()["type"] == "urn:crm:problem:unsupported-format"

    pdf = await client.get(f"{url}?tool=canva&format=pdf")
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    png = await client.get(f"{url}?tool=figma&format=png")
    assert png.headers["content-type"] == "application/zip"
    names = zipfile.ZipFile(io.BytesIO(png.content)).namelist()
    assert names == [f"{demand['demand_code']}_v2_page-1.png", f"{demand['demand_code']}_v2_page-2.png"]
    assert len(calls) == 1  # the PDF is converted once per version, then reused

    # Edits after approval: the export follows the newest version.
    await send(client, demand["id"], "make the title shorter")
    await work(settings)
    newest = await client.get(f"{url}?tool=canva&format=pdf")
    assert f"{demand['demand_code']}_v3.pdf" in newest.headers["content-disposition"] and len(calls) == 2


async def test_export_without_libreoffice_says_so(client, make_user, admin_password_set, ai, settings, monkeypatch):
    demand = await approved_letter(client, make_user, admin_password_set, ai, settings)
    monkeypatch.setattr("app.modules.ads.convert.find_soffice", lambda configured: None)
    response = await client.get(f"/api/v1/ads/{demand['id']}/export?tool=canva&format=pdf")
    assert response.status_code == 503 and response.json()["type"] == "urn:crm:problem:converter-missing"
