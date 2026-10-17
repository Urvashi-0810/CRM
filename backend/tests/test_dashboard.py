"""Dashboard: periods (India time), orders/vacancies/clients, repeat clients, countries, ad status."""

import io
from datetime import date, datetime
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook

from app.errors import ProblemError
from app.modules.dashboard.service import previous_range, resolve_period
from tests.conftest import login
from tests.test_demands import add, approve_directly

IST = ZoneInfo("Asia/Kolkata")


def ist(y, m, d, hh=12, mm=0):
    return datetime(y, m, d, hh, mm, tzinfo=IST)


# ---------- periods ----------


def test_periods():
    today = date(2026, 3, 15)
    r = resolve_period("this_month", None, None, today=today)
    assert (r.start, r.end, r.label) == (date(2026, 3, 1), today, "March 2026")
    r = resolve_period("last_month", None, None, today=today)
    assert (r.start, r.end) == (date(2026, 2, 1), date(2026, 2, 28))
    r = resolve_period("last_month", None, None, today=date(2026, 1, 10))
    assert (r.start, r.end, r.label) == (date(2025, 12, 1), date(2025, 12, 31), "December 2025")
    r = resolve_period("this_year", None, None, today=today)
    assert (r.start, r.end) == (date(2026, 1, 1), today)
    r = resolve_period("custom", date(2026, 1, 5), date(2026, 2, 5))
    assert (r.start, r.end) == (date(2026, 1, 5), date(2026, 2, 5))
    with pytest.raises(ProblemError):
        resolve_period("custom", date(2026, 2, 5), date(2026, 1, 5))
    with pytest.raises(ProblemError):
        resolve_period("custom", None, date(2026, 1, 5))


def test_comparison_periods():
    # This month so far → the same days of last month (cut at its last day).
    p = previous_range(resolve_period("this_month", None, None, today=date(2026, 3, 31)))
    assert (p.start, p.end, p.label) == (date(2026, 2, 1), date(2026, 2, 28), "1–28 Feb")
    p = previous_range(resolve_period("last_month", None, None, today=date(2026, 3, 5)))
    assert (p.start, p.end, p.label) == (date(2026, 1, 1), date(2026, 1, 31), "January")
    p = previous_range(resolve_period("this_year", None, None, today=date(2028, 2, 29)))
    assert (p.start, p.end, p.label) == (date(2027, 1, 1), date(2027, 2, 28), "same time 2027")
    p = previous_range(resolve_period("custom", date(2026, 3, 1), date(2026, 3, 10)))
    assert (p.start, p.end, p.label) == (date(2026, 2, 19), date(2026, 2, 28), "previous 10 days")


# ---------- numbers ----------


async def _approved(client, user_id, when, **fields):
    demand_id = (await add(client)).json()["id"]
    await approve_directly(demand_id, user_id, when, **fields)
    return demand_id


async def test_dashboard_counts(client, make_user):
    user = await make_user("priya", "Priya Shah")
    await login(client, "priya")
    # March: Mosanada twice (one with odd spacing/case) → repeat client; Axion once.
    await _approved(client, user.id, ist(2026, 3, 2), employer="Mosanada", country="Qatar", vacancy=10)
    await _approved(client, user.id, ist(2026, 3, 20), employer=" mosanada ", country="qatar", vacancy=5)
    await _approved(client, user.id, ist(2026, 3, 25), employer="Axion", country="UAE", vacancy=40)
    # The AI's placeholder is not a client.
    await _approved(client, user.id, ist(2026, 3, 26), employer="[[MISSING: Employer Name]]", vacancy=None)
    # 1 April 00:30 India time is still 31 March in UTC: it belongs to April.
    await _approved(client, user.id, ist(2026, 4, 1, 0, 30), employer="Axion", country="UAE", vacancy=3)
    # Drafts are never counted.
    draft = (await add(client)).json()["id"]
    assert draft

    march = (await client.get("/api/v1/dashboard?period=custom&date_from=2026-03-01&date_to=2026-03-31")).json()
    t = march["totals"]
    assert (t["orders"], t["vacancies"], t["clients"], t["new_clients"]) == (4, 55, 2, 2)
    assert march["previous_totals"]["orders"] == 0 and march["previous"]["label"] == "previous 31 days"
    assert march["trend"]["unit"] == "day" and len(march["trend"]["points"]) == 31
    assert [p["bucket"] for p in march["trend"]["points"] if p["orders"]] == ["2026-03-02", "2026-03-20", "2026-03-25", "2026-03-26"]
    assert march["team"] == [{"name": "Priya Shah", "orders": 4, "vacancies": 55}]
    assert march["positions"] == []
    assert (march["pipeline"]["drafts"], march["pipeline"]["approved_all_time"]) == (1, 5)
    assert [(c["employer"], c["orders"], c["vacancies"]) for c in march["by_client"]] == [
        ("mosanada", 2, 15), ("Axion", 1, 40), ("Not filled", 1, 0)]
    assert march["repeat_clients"] == [{"employer": "mosanada", "repeat_months": 1, "orders_in_repeat_months": 2,
                                         "months": [{"month": "2026-03", "orders": 2}]}]
    assert [(c["country"], c["orders"]) for c in march["by_country"]] == [("qatar", 2), ("UAE", 1), ("Not filled", 1)]
    assert march["ad_status"] == [{"status": "Not started", "orders": 4}, {"status": "In progress", "orders": 0},
                                  {"status": "Done", "orders": 0}]

    april = (await client.get("/api/v1/dashboard?period=custom&date_from=2026-04-01&date_to=2026-04-30")).json()
    t = april["totals"]
    # Axion's first order was in March: a client in April, but not a new one.
    assert (t["orders"], t["vacancies"], t["clients"], t["new_clients"]) == (1, 3, 1, 0)
    assert april["repeat_clients"] == [] and april["previous_totals"]["orders"] == 4  # 2–31 March

    # A longer period: repeat = any single month with 2+ orders; Axion's two orders are in different months.
    spring = (await client.get("/api/v1/dashboard?period=custom&date_from=2026-03-01&date_to=2026-04-30")).json()
    assert [c["employer"] for c in spring["repeat_clients"]] == ["mosanada"]
    assert next(c for c in spring["by_client"] if c["employer"] == "Axion")["orders"] == 2

    bad = await client.get("/api/v1/dashboard?period=custom&date_from=2026-04-30&date_to=2026-03-01")
    assert bad.status_code == 422

    export = await client.get("/api/v1/dashboard/export?period=custom&date_from=2026-03-01&date_to=2026-03-31")
    assert export.status_code == 200 and "dashboard_2026-03-01_to_2026-03-31.xlsx" in export.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(export.content))
    assert wb.sheetnames == ["Summary", "Orders over time", "Orders by client", "Repeat clients", "Orders by country",
                             "Top positions", "Ad status", "Team"]
    assert ["Orders (approved letters)", 4, 0] in [list(r) for r in wb["Summary"].iter_rows(values_only=True)]


async def test_dashboard_needs_login(client):
    assert (await client.get("/api/v1/dashboard")).status_code == 401
