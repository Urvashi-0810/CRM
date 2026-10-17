import io
import zipfile
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from docx import Document
from PIL import Image

from app.modules.extraction.readers import UploadRejected, read_upload
from app.modules.letters.docx_builder import BuildInput, LetterheadImages, build_docx
from app.modules.letters.letter_model import AiResponse, Letter
from app.modules.llm.pricing import cost_usd

FIX = Path(__file__).parent / "fixtures"


def png(w=200, h=100, color=(10, 120, 200)) -> bytes:
    out = io.BytesIO()
    Image.new("RGB", (w, h), color).save(out, "PNG")
    return out.getvalue()


# ---------- readers on the client's real samples ----------


def test_inquiry_with_multiple_project_tables():
    r = read_upload((FIX / "inquiry1_mosanada.docx").read_bytes(), "inquiry.docx", keep_images=False)
    assert r.mode == "text"
    assert "Project: Aspire Zone" in r.text and "Project: ISF" in r.text
    assert "| HVAC/Mechanical Engineer | 2 | QAR 6000 + Airfare + Acc & transport |" in r.text


def test_inquiry_shorthand_kept_exactly():
    r = read_upload((FIX / "inquiry4_hvac.docx").read_bytes(), "i.docx", keep_images=False)
    assert "4000+FAT" in r.text  # shorthand stays as written


def test_jd_excel_list():
    r = read_upload((FIX / "jd_list_of_resources.xlsx").read_bytes(), "jd.xlsx", keep_images=True)
    assert "Sheet: List" in r.text
    assert "Genset Technician" in r.text and "| 16 |" in r.text


def test_docx_pictures_are_extracted_for_jd():
    doc = Document()
    doc.add_paragraph("Site photo below")
    doc.add_picture(io.BytesIO(png(300, 200)))
    doc.add_picture(io.BytesIO(png(20, 20)))  # icon-sized: ignored
    buf = io.BytesIO()
    doc.save(buf)
    r = read_upload(buf.getvalue(), "jd.docx", keep_images=True)
    assert [(i.width, i.height) for i in r.images] == [(300, 200)]
    assert read_upload(buf.getvalue(), "inq.docx", keep_images=False).images == []


def test_old_word_format_rejected_with_save_as_hint():
    with pytest.raises(UploadRejected) as exc:
        read_upload(b"\xd0\xcf\x11\xe0", "old.doc", keep_images=False)
    assert exc.value.status == 415 and ".docx" in exc.value.detail


def test_corrupt_docx_rejected():
    with pytest.raises(UploadRejected) as exc:
        read_upload(b"not a zip", "x.docx", keep_images=False)
    assert exc.value.slug == "file-unreadable"


def test_image_too_large_in_pixels():
    big = io.BytesIO()
    Image.new("RGB", (8100, 10)).save(big, "PNG")
    with pytest.raises(UploadRejected) as exc:
        read_upload(big.getvalue(), "big.png", keep_images=False)
    assert exc.value.status == 413 and exc.value.title == "File too large"


def test_pdf_page_limit():
    from pypdf import PdfWriter

    w = PdfWriter()
    for _ in range(101):
        w.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    w.write(buf)
    with pytest.raises(UploadRejected) as exc:
        read_upload(buf.getvalue(), "long.pdf", keep_images=False)
    assert exc.value.status == 413
    ok = PdfWriter()
    ok.add_blank_page(width=200, height=200)
    small = io.BytesIO()
    ok.write(small)
    assert read_upload(small.getvalue(), "short.pdf", keep_images=False).pages == 1


# ---------- letter model ----------


def test_letter_normalising_drops_unknown_images_and_pads_tables():
    letter = Letter.model_validate({"title": "T", "blocks": [
        {"type": "paragraph", "text": "  "},
        {"type": "image", "image_id": "img-99"},
        {"type": "image", "image_id": "img-5"},
        {"type": "table", "columns": ["A", "B", "C"], "rows": [["1"], ["", "", ""]]},
    ]}).normalised({"img-5"})
    assert [b.type for b in letter.blocks] == ["image", "table"]
    assert letter.blocks[1].rows == [["1", "", ""]]


def test_missing_markers_collected():
    letter = Letter.model_validate({"blocks": [
        {"type": "paragraph", "text": "Employer: [[MISSING: Employer]]"},
        {"type": "table", "columns": ["Salary"], "rows": [["[[MISSING: Currency]]"]]},
    ]})
    assert letter.missing_markers() == ["Employer", "Currency"]


def test_ai_response_parses():
    r = AiResponse.model_validate_json(
        '{"action":"ask","message":"Need details","questions":["Which country?"],'
        '"missing_fields":[],"letter":null,"extracted":null}'
    )
    assert r.action == "ask" and r.questions == ["Which country?"]


# ---------- docx ----------


def _docx_text(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return " ".join(z.read(n).decode("utf8", "ignore") for n in z.namelist() if n.startswith("word/"))


def test_docx_has_letterhead_id_draft_mark_and_highlight():
    letter = Letter.model_validate({"title": "Demand Letter", "blocks": [
        {"type": "heading", "text": "Positions", "level": 2},
        {"type": "table", "columns": ["Position", "Qty"], "rows": [["Welder", "4"]]},
        {"type": "bullets", "items": ["Contract: [[MISSING: Contract period]]"]},
        {"type": "image", "image_id": "img-1", "caption": "Site"},
    ]})
    data = build_docx(BuildInput(
        letter=letter, demand_code="AA00042", approval_date=None, is_draft=True,
        letterhead=LetterheadImages(
            (FIX / "test_header.png").read_bytes(), (FIX / "test_footer.png").read_bytes()),
        images={"img-1": png()},
    ))
    xml = _docx_text(data)
    assert "AA00042" in xml and "DRAFT — NOT APPROVED" in xml
    assert "Contract period — to be filled" in xml and 'w:highlight w:val="yellow"' in xml
    doc = Document(io.BytesIO(data))
    assert len(doc.inline_shapes) == 1  # the JD image in the body
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        media = [n for n in z.namelist() if n.startswith("word/media/")]
    assert len(media) == 3  # header, footer, body image


def test_docx_approved_shows_date_and_no_draft_mark():
    data = build_docx(BuildInput(
        letter=Letter(blocks=[]), demand_code="AA00001", approval_date=date(2026, 10, 14),
        is_draft=False, letterhead=None, images={},
    ))
    xml = _docx_text(data)
    assert "14 October 2026" in xml and "DRAFT" not in xml


# ---------- pricing ----------


def test_claude_cost():
    # Sonnet 5.5: $2 in / $10 out / $0.10 cache read / $2.50 cache write per 1M
    assert cost_usd("claude", "claude-sonnet-5-5", input_tokens=1_000_000, output_tokens=100_000,
                     cache_read_tokens=1_000_000) == Decimal("3.100000")


def test_haiku_long_prompt_tier_and_unknown_model():
    assert cost_usd("claude", "claude-haiku-5-5", input_tokens=200_000, output_tokens=0) == Decimal("0.100000")
    assert cost_usd("claude", "claude-unknown-1", input_tokens=10, output_tokens=10) is None


def test_gemini_flash_lite_not_priced_as_flash():
    assert cost_usd("gemini", "gemini-3.5-flash-lite", input_tokens=1_000_000, output_tokens=0) == Decimal("0.300000")
