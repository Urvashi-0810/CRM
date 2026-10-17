import asyncio
import io
import uuid
from datetime import UTC, datetime, timedelta

from openpyxl import load_workbook
from sqlalchemy import text

from app.db import get_sessionmaker
from app.modules.demands.models import format_demand_code
from app.modules.storage import StorageError, get_storage
from app.modules.storage.local import LocalStorage
from tests.conftest import login


async def add(client, key: str | None = None):
    headers = {"Idempotency-Key": key} if key else {}
    return await client.post("/api/v1/demands", headers=headers)


async def approve_directly(demand_id: int, user_id: int, when: datetime, **fields):
    """Approval arrives in Milestone 4; tests set the row directly."""
    sets = ", ".join(f"{k} = :{k}" for k in fields)
    async with get_sessionmaker()() as db:
        await db.execute(
            text(
                "update demand_letters set lifecycle = 'approved', approved_by = :u, approved_at = :t"
                + (f", {sets}" if sets else "")
                + " where id = :id"
            ),
            {"u": user_id, "t": when, "id": demand_id, **fields},
        )
        await db.commit()


# ---------- codes ----------


def test_demand_code_format():
    assert format_demand_code(1) == "AA00001"
    assert format_demand_code(99_999) == "AA99999"
    assert format_demand_code(100_000) == "AB00001"
    assert format_demand_code(26 * 99_999 + 1) == "BA00001"
    assert format_demand_code(26 * 26 * 99_999) == "ZZ99999"


# ---------- creating ----------


async def test_add_demand_letter_creates_draft_and_folder(client, make_user, storage_root):
    await make_user("priya", "Priya Shah")
    await login(client, "priya")

    response = await add(client)

    assert response.status_code == 201
    body = response.json()
    assert body["demand_code"] == "AA00001"
    assert body["lifecycle"] == "draft"
    assert body["progress_status"] == "Start"
    assert body["created_by"]["display_name"] == "Priya Shah"
    assert response.headers["etag"] == body["etag"]
    assert response.headers["location"].endswith(f"/demands/{body['id']}")
    assert (storage_root / "demands" / "AA00001").is_dir()


async def test_twenty_simultaneous_adds_get_twenty_different_lowest_ids(client, new_client, make_user):
    await make_user("priya")
    browsers = []
    for _ in range(20):
        c = await new_client()
        await login(c, "priya")
        browsers.append(c)

    responses = await asyncio.gather(*(add(c) for c in browsers))

    assert all(r.status_code == 201 for r in responses)
    codes = sorted(r.json()["demand_code"] for r in responses)
    assert codes == [format_demand_code(n) for n in range(1, 21)]


async def test_deleted_draft_id_is_reused_lowest_first(client, make_user, storage_root):
    await make_user("priya")
    await login(client, "priya")
    ids = [(await add(client)).json()["id"] for _ in range(5)]  # AA00001..AA00005

    assert (await client.delete(f"/api/v1/demands/{ids[3]}")).status_code == 204  # AA00004
    assert (await client.delete(f"/api/v1/demands/{ids[1]}")).status_code == 204  # AA00002
    assert not (storage_root / "demands" / "AA00002").exists()

    assert (await add(client)).json()["demand_code"] == "AA00002"
    assert (await add(client)).json()["demand_code"] == "AA00004"
    assert (await add(client)).json()["demand_code"] == "AA00006"


async def test_same_idempotency_key_returns_same_draft(client, make_user):
    await make_user("priya")
    await login(client, "priya")
    key = str(uuid.uuid4())

    first, second = await asyncio.gather(add(client, key), add(client, key))

    assert {first.status_code, second.status_code} == {200, 201}
    assert first.json()["id"] == second.json()["id"]
    page = (await client.get("/api/v1/demands?lifecycle=draft")).json()
    assert page["total"] == 1


async def test_idempotency_key_after_draft_deleted(client, make_user):
    await make_user("priya")
    await login(client, "priya")
    key = str(uuid.uuid4())
    demand = (await add(client, key)).json()
    await client.delete(f"/api/v1/demands/{demand['id']}")

    again = await add(client, key)
    assert again.status_code == 409


# ---------- editing ----------


async def test_edit_fields_with_etag(client, make_user):
    await make_user("priya")
    await login(client, "priya")
    created = await add(client)
    etag = created.headers["etag"]
    demand_id = created.json()["id"]

    updated = await client.patch(
        f"/api/v1/demands/{demand_id}",
        json={"country": " Qatar ", "employer": "Mosanada FMS", "vacancy": 25,
              "target_timeline": "45 days", "assign_to": "Anita"},
        headers={"If-Match": etag},
    )
    assert updated.status_code == 200
    body = updated.json()
    assert body["country"] == "Qatar"
    assert body["vacancy"] == 25
    assert updated.headers["etag"] != etag

    # Saving with the old ETag (someone else changed it meanwhile) is refused.
    stale = await client.patch(
        f"/api/v1/demands/{demand_id}", json={"assign_to": "Ravi"}, headers={"If-Match": etag}
    )
    assert stale.status_code == 412
    assert stale.json()["type"] == "urn:crm:problem:demand-changed"

    # Clearing a field with "" works; untouched fields stay.
    cleared = await client.patch(
        f"/api/v1/demands/{demand_id}", json={"assign_to": ""},
        headers={"If-Match": updated.headers["etag"]},
    )
    assert cleared.json()["assign_to"] is None
    assert cleared.json()["employer"] == "Mosanada FMS"


async def test_edit_without_if_match_is_428(client, make_user):
    await make_user("priya")
    await login(client, "priya")
    demand_id = (await add(client)).json()["id"]
    response = await client.patch(f"/api/v1/demands/{demand_id}", json={"country": "UAE"})
    assert response.status_code == 428


async def test_invalid_vacancy_rejected(client, make_user):
    await make_user("priya")
    await login(client, "priya")
    created = await add(client)
    response = await client.patch(
        f"/api/v1/demands/{created.json()['id']}", json={"vacancy": -3},
        headers={"If-Match": created.headers["etag"]},
    )
    assert response.status_code == 422


async def test_replace_trades(client, make_user):
    await make_user("priya")
    await login(client, "priya")
    created = await add(client)
    demand_id = created.json()["id"]

    response = await client.put(
        f"/api/v1/demands/{demand_id}/trades",
        json=[
            {"project": "Aspire Zone", "trade": "HVAC/Mechanical Engineer", "quantity": 2},
            {"project": "", "trade": "Housekeeping Officer", "quantity": 10},
        ],
        headers={"If-Match": created.headers["etag"]},
    )
    assert response.status_code == 200
    trades = response.json()["trades"]
    assert [t["trade"] for t in trades] == ["HVAC/Mechanical Engineer", "Housekeeping Officer"]
    assert trades[1]["project"] is None

    bad = await client.put(
        f"/api/v1/demands/{demand_id}/trades",
        json=[{"trade": "Welder", "quantity": 0}],
        headers={"If-Match": response.headers["etag"]},
    )
    assert bad.status_code == 422

    emptied = await client.put(
        f"/api/v1/demands/{demand_id}/trades", json=[], headers={"If-Match": response.headers["etag"]}
    )
    assert emptied.json()["trades"] == []


# ---------- deleting ----------


async def test_approved_letter_cannot_be_deleted(client, make_user):
    user = await make_user("priya")
    await login(client, "priya")
    demand_id = (await add(client)).json()["id"]
    await approve_directly(demand_id, user.id, datetime.now(UTC))

    response = await client.delete(f"/api/v1/demands/{demand_id}")
    assert response.status_code == 409
    assert response.json()["type"] == "urn:crm:problem:approved-letter-not-deletable"


async def test_folder_delete_failure_keeps_draft_and_id(app, client, make_user, settings):
    await make_user("priya")
    await login(client, "priya")
    demand_id = (await add(client)).json()["id"]

    class BrokenDelete(LocalStorage):
        async def delete_folder(self, prefix):
            raise StorageError("disk error")

    app.dependency_overrides[get_storage] = lambda: BrokenDelete(settings.storage_root)
    response = await client.delete(f"/api/v1/demands/{demand_id}")
    assert response.status_code == 503
    assert response.json()["type"] == "urn:crm:problem:storage-delete-failed"
    assert (await client.get(f"/api/v1/demands/{demand_id}")).status_code == 200

    app.dependency_overrides.pop(get_storage)
    assert (await add(client)).json()["demand_code"] == "AA00002"  # AA00001 still reserved


async def test_delete_again_after_folder_already_gone(client, make_user, storage_root):
    await make_user("priya")
    await login(client, "priya")
    demand_id = (await add(client)).json()["id"]
    (storage_root / "demands" / "AA00001").rmdir()  # e.g. a previous attempt removed it

    assert (await client.delete(f"/api/v1/demands/{demand_id}")).status_code == 204
    assert (await client.get(f"/api/v1/demands/{demand_id}")).status_code == 404


# ---------- list, filters, export ----------


async def _seed_list(client, user_id):
    """Three approved letters + one draft."""
    now = datetime.now(UTC)
    rows = [
        dict(country="Qatar", employer="Mosanada FMS", vacancy=25, assign_to="Anita", when=now - timedelta(days=2)),
        dict(country="UAE", employer="Axion Prime Contracting", vacancy=67, assign_to="Ravi", when=now - timedelta(days=10)),
        dict(country="Qatar", employer="AQI", vacancy=2, assign_to="Anita", when=now - timedelta(days=40)),
    ]
    ids = []
    for row in rows:
        demand_id = (await add(client)).json()["id"]
        when = row.pop("when")
        await approve_directly(demand_id, user_id, when, **row)
        ids.append(demand_id)
    await add(client)  # stays a draft
    return ids


async def test_list_filters_sort_and_pagination(client, make_user):
    user = await make_user("priya", "Priya Shah")
    await login(client, "priya")
    await _seed_list(client, user.id)

    page = (await client.get("/api/v1/demands")).json()
    assert page["total"] == 3
    assert [i["employer"] for i in page["items"]] == [
        "Mosanada FMS", "Axion Prime Contracting", "AQI"]  # newest approval first
    assert page["items"][0]["approved_by"]["display_name"] == "Priya Shah"

    drafts = (await client.get("/api/v1/demands?lifecycle=draft")).json()
    assert drafts["total"] == 1 and drafts["items"][0]["demand_code"] == "AA00004"

    qatar = (await client.get("/api/v1/demands?country=Qatar")).json()
    assert qatar["total"] == 2
    search = (await client.get("/api/v1/demands?q=axion")).json()
    assert search["total"] == 1
    by_code = (await client.get("/api/v1/demands?q=aa00003")).json()
    assert by_code["items"][0]["employer"] == "AQI"

    paged = (await client.get("/api/v1/demands?limit=2&offset=2")).json()
    assert paged["total"] == 3 and len(paged["items"]) == 1

    vacancy_sorted = (await client.get("/api/v1/demands?sort=-vacancy")).json()
    assert [i["vacancy"] for i in vacancy_sorted["items"]] == [67, 25, 2]

    today = datetime.now(UTC).date()
    recent = (await client.get(
        f"/api/v1/demands?approved_from={(today - timedelta(days=15)).isoformat()}"
        f"&approved_to={today.isoformat()}"
    )).json()
    assert recent["total"] == 2

    backwards = await client.get("/api/v1/demands?approved_from=2026-10-09&approved_to=2026-10-01")
    assert backwards.status_code == 422

    options = (await client.get("/api/v1/demands/filter-options")).json()
    assert options["countries"] == ["Qatar", "UAE"]
    assert options["assignees"] == ["Anita", "Ravi"]
    assert options["statuses"] == ["Start"]


async def test_approved_by_shows_current_name(client, make_user):
    user = await make_user("priya", "Priya Shah")
    await login(client, "priya")
    demand_id = (await add(client)).json()["id"]
    await approve_directly(demand_id, user.id, datetime.now(UTC))
    async with get_sessionmaker()() as db:
        await db.execute(text("update users set display_name = 'Priya Mehta' where id = :id"), {"id": user.id})
        await db.commit()

    item = (await client.get("/api/v1/demands")).json()["items"][0]
    assert item["approved_by"]["display_name"] == "Priya Mehta"


async def test_export_xlsx_follows_filters(client, make_user):
    user = await make_user("priya", "Priya Shah")
    await login(client, "priya")
    await _seed_list(client, user.id)
    # Text starting with "=" must not become a formula in Excel.
    demand_id = (await add(client)).json()["id"]
    await approve_directly(demand_id, user.id, datetime.now(UTC), country="Qatar", employer="=HYPERLINK(1)")

    response = await client.get("/api/v1/demands/export?country=Qatar")

    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]
    assert response.headers["content-disposition"].startswith('attachment; filename="demand-letters_')
    sheet = load_workbook(io.BytesIO(response.content)).active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0][:4] == ("Approved date", "Country", "Demand ID", "Employer")
    assert rows[0][-2:] == ("Created by", "Approved by")
    assert len(rows) == 1 + 3
    assert {r[1] for r in rows[1:]} == {"Qatar"}
    assert "'=HYPERLINK(1)" in {r[3] for r in rows[1:]}


async def test_demands_require_login(client):
    assert (await client.get("/api/v1/demands")).status_code == 401
    assert (await client.post("/api/v1/demands")).status_code == 401


async def test_unknown_demand_is_404(client, make_user):
    await make_user("priya")
    await login(client, "priya")
    assert (await client.get("/api/v1/demands/999")).status_code == 404
    assert (await client.delete("/api/v1/demands/999")).status_code == 404
