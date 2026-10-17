"""Approve, versions, restore, upload an edited letter, reference letters, spend cap, usage."""

import io
import json
import zipfile
from decimal import Decimal
from pathlib import Path

from docx import Document
from sqlalchemy import text

from app.db import get_sessionmaker
from tests.conftest import login, unlock_admin
from tests.test_chat_flow import (
    ASK,
    DOCX,
    LETTER,
    ai,  # noqa: F401  (fixture)
    inquiry_file,
    send,
    setup,
    work,
)

FIX = Path(__file__).parent / "fixtures"


def header_xml(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return " ".join(z.read(n).decode() for n in z.namelist() if n.startswith("word/header"))


async def letter_ready(client, make_user, admin_password_set, ai, settings, replies=(LETTER,)):
    ai(*replies)
    demand = await setup(client, make_user, admin_password_set)
    await send(client, demand["id"], "", [inquiry_file()])
    await work(settings)
    return demand


async def submit_and_approve(client, demand_id, version_no):
    """Send for approval, then approve as the admin (the test client has the admin unlocked)."""
    submitted = await client.post(f"/api/v1/demands/{demand_id}/submit", json={"version_no": version_no})
    assert submitted.status_code == 204, submitted.text
    return await client.post(f"/api/v1/demands/{demand_id}/approve", json={"version_no": version_no})


# ---------------- send for approval / approve ----------------


async def test_send_for_approval_then_admin_approves_or_sends_back(
    client, make_user, admin_password_set, ai, settings
):
    demand = await letter_ready(client, make_user, admin_password_set, ai, settings)
    url = f"/api/v1/demands/{demand['id']}"

    # A draft can't be approved before it is sent.
    early = await client.post(f"{url}/approve", json={"version_no": 1})
    assert early.status_code == 409 and early.json()["type"] == "urn:crm:problem:not-pending"

    assert (await client.post(f"{url}/submit", json={"version_no": 1})).status_code == 204
    detail = (await client.get(url)).json()
    assert detail["lifecycle"] == "pending" and detail["submitted_by"]["display_name"] == "Priya Shah"
    assert detail["submitted_at"] and detail["approved_by"] is None
    pending = (await client.get("/api/v1/demands?lifecycle=pending")).json()
    assert [i["demand_code"] for i in pending["items"]] == [demand["demand_code"]]
    assert (await client.get("/api/v1/demands?lifecycle=draft")).json()["total"] == 0
    assert (await client.get("/api/v1/demands")).json()["total"] == 0  # not in the list until approved
    # Still a draft on paper while it waits.
    assert "DRAFT — NOT APPROVED" in header_xml((await client.get(f"{url}/docx")).content)
    again = await client.post(f"{url}/submit", json={"version_no": 1})
    assert again.status_code == 409 and again.json()["type"] == "urn:crm:problem:already-submitted"

    # Only an unlocked admin can approve or send back.
    assert (await client.delete("/api/v1/admin/session")).status_code == 204
    for path, body in (("approve", {"version_no": 1}), ("send-back", {})):
        refused = await client.post(f"{url}/{path}", json=body)
        assert refused.status_code == 403 and refused.json()["type"] == "urn:crm:problem:admin-session-required"
    assert (await unlock_admin(client)).status_code == 204

    back = await client.post(f"{url}/send-back", json={"reason": "Salary is missing"})
    assert back.status_code == 204
    detail = (await client.get(url)).json()
    assert detail["lifecycle"] == "draft" and detail["submitted_by"] is None
    events = [m for m in (await client.get(f"{url}/messages")).json()["items"] if m["kind"] == "event"]
    assert "Sent back to draft by Priya Shah. Reason: Salary is missing" in events[-1]["body"]
    assert (await client.post(f"{url}/send-back", json={})).status_code == 409

    assert (await submit_and_approve(client, demand["id"], 1)).status_code == 200
    assert (await client.get(url)).json()["lifecycle"] == "approved"



async def test_approve_stamps_date_removes_draft_and_moves_to_list(client, make_user, admin_password_set, ai, settings):
    demand = await letter_ready(client, make_user, admin_password_set, ai, settings)
    before = (await client.get(f"/api/v1/demands/{demand['id']}/docx")).content
    assert "DRAFT — NOT APPROVED" in header_xml(before)

    response = await submit_and_approve(client, demand["id"], 1)
    assert response.status_code == 200, response.text
    assert response.json() == {"version_no": 2, "rebuilt": False}

    detail = (await client.get(f"/api/v1/demands/{demand['id']}")).json()
    assert detail["lifecycle"] == "approved" and detail["approved_by"]["display_name"] == "Priya Shah"
    after = (await client.get(f"/api/v1/demands/{demand['id']}/docx")).content
    xml = header_xml(after)
    assert "DRAFT" not in xml and "Date :" in xml
    from datetime import datetime
    from zoneinfo import ZoneInfo

    assert datetime.now(ZoneInfo("Asia/Kolkata")).strftime("%d %B %Y") in xml
    # The body is untouched by stamping.
    assert Document(io.BytesIO(after)).tables[0]._cells[-1].text == Document(io.BytesIO(before)).tables[0]._cells[-1].text

    listing = (await client.get("/api/v1/demands")).json()
    assert [i["demand_code"] for i in listing["items"]] == [demand["demand_code"]]
    assert (await client.get("/api/v1/demands?lifecycle=draft")).json()["total"] == 0
    events = [m for m in (await client.get(f"/api/v1/demands/{demand['id']}/messages")).json()["items"] if m["kind"] == "event"]
    assert "Approved by Priya Shah" in events[-1]["body"]

    again = await client.post(f"/api/v1/demands/{demand['id']}/approve", json={"version_no": 2})
    assert again.status_code == 409 and again.json()["type"] == "urn:crm:problem:already-approved"
    assert (await client.delete(f"/api/v1/demands/{demand['id']}")).status_code == 409


async def test_approve_rules(client, make_user, admin_password_set, ai, settings):
    ai(ASK)
    demand = await setup(client, make_user, admin_password_set)
    none_yet = await client.post(f"/api/v1/demands/{demand['id']}/submit", json={"version_no": 1})
    assert none_yet.status_code == 422

    ai(LETTER, LETTER, LETTER)
    await send(client, demand["id"], "", [inquiry_file()])
    busy = await client.post(f"/api/v1/demands/{demand['id']}/submit", json={"version_no": 1})
    assert busy.status_code == 409 and busy.json()["type"] == "urn:crm:problem:job-in-progress"
    await work(settings)
    await send(client, demand["id"], "change something")
    await work(settings)
    stale = await client.post(f"/api/v1/demands/{demand['id']}/submit", json={"version_no": 1})
    assert stale.status_code == 409 and stale.json()["type"] == "urn:crm:problem:stale-version"
    # The admin also approves only the newest version.
    assert (await client.post(f"/api/v1/demands/{demand['id']}/submit", json={"version_no": 2})).status_code == 204
    await send(client, demand["id"], "one more change")
    await work(settings)
    stale = await client.post(f"/api/v1/demands/{demand['id']}/approve", json={"version_no": 2})
    assert stale.status_code == 409 and stale.json()["type"] == "urn:crm:problem:stale-version"


async def test_prompt_edit_after_approval_keeps_date(client, make_user, admin_password_set, ai, settings):
    demand = await letter_ready(client, make_user, admin_password_set, ai, settings, replies=(LETTER, LETTER))
    await submit_and_approve(client, demand["id"], 1)
    approved_at = (await client.get(f"/api/v1/demands/{demand['id']}")).json()["approved_at"]

    await send(client, demand["id"], "make the title shorter")
    await work(settings)
    letter = (await client.get(f"/api/v1/demands/{demand['id']}/letter")).json()
    assert letter["version_no"] == 3 and letter["source"] == "prompt_edit"
    detail = (await client.get(f"/api/v1/demands/{demand['id']}")).json()
    assert detail["approved_at"] == approved_at and detail["lifecycle"] == "approved"
    xml = header_xml((await client.get(f"/api/v1/demands/{demand['id']}/docx")).content)
    assert "DRAFT" not in xml


# ---------------- versions + restore ----------------


async def test_version_history_and_restore(client, make_user, admin_password_set, ai, settings):
    edited = json.loads(LETTER)
    edited["letter"]["title"] = "Changed title"
    demand = await letter_ready(client, make_user, admin_password_set, ai, settings,
                                replies=(LETTER, json.dumps(edited)))
    await send(client, demand["id"], "change the title")
    await work(settings)

    versions = (await client.get(f"/api/v1/demands/{demand['id']}/versions")).json()["items"]
    assert [(v["version_no"], v["source"]) for v in versions] == [(2, "prompt_edit"), (1, "generated")]
    v1 = (await client.get(f"/api/v1/demands/{demand['id']}/versions/1")).json()
    assert v1["content"]["title"] == "Demand Letter"
    old_file = await client.get(f"/api/v1/demands/{demand['id']}/versions/1/docx")
    assert old_file.status_code == 200 and "_v1.docx" in old_file.headers["content-disposition"]

    restored = await client.post(f"/api/v1/demands/{demand['id']}/versions/1/restore")
    assert restored.status_code == 201 and restored.json() == {"version_no": 3}
    latest = (await client.get(f"/api/v1/demands/{demand['id']}/letter")).json()
    assert latest["content"]["title"] == "Demand Letter" and latest["source"] == "restored"
    assert (await client.get(f"/api/v1/demands/{demand['id']}/versions")).json()["items"][0]["restored_from_no"] == 1
    assert (await client.post(f"/api/v1/demands/{demand['id']}/versions/3/restore")).status_code == 409
    assert (await client.get(f"/api/v1/demands/{demand['id']}/versions/9")).status_code == 404


async def test_restoring_draft_version_of_approved_letter_gets_approval_stamp(
    client, make_user, admin_password_set, ai, settings
):
    demand = await letter_ready(client, make_user, admin_password_set, ai, settings)
    await submit_and_approve(client, demand["id"], 1)
    await client.post(f"/api/v1/demands/{demand['id']}/versions/1/restore")  # v1 was a draft
    xml = header_xml((await client.get(f"/api/v1/demands/{demand['id']}/docx")).content)
    assert "DRAFT" not in xml


# ---------------- upload an edited letter ----------------


async def test_upload_edited_letter_round_trip(client, make_user, admin_password_set, ai, settings):
    adapter = ai(LETTER, LETTER)
    demand = await letter_ready(client, make_user, admin_password_set, ai, settings, replies=(LETTER, LETTER))
    adapter = ai(LETTER)  # for the prompt edit after the upload

    original = (await client.get(f"/api/v1/demands/{demand['id']}/docx")).content
    doc = Document(io.BytesIO(original))
    doc.tables[0].rows[1].cells[2].text = "3"            # Aspire Zone engineers: 2 → 3
    doc.add_paragraph("Added by hand in Word.")
    buf = io.BytesIO()
    doc.save(buf)

    up = await client.post(f"/api/v1/demands/{demand['id']}/edited-uploads",
                           files={"file": ("edited.docx", buf.getvalue(), DOCX)})
    assert up.status_code == 201, up.text
    assert up.json() == {"version_no": 2, "header_found": True}

    letter = (await client.get(f"/api/v1/demands/{demand['id']}/letter")).json()
    assert letter["source"] == "uploaded_edit"
    v1 = (await client.get(f"/api/v1/demands/{demand['id']}/versions/1")).json()
    assert letter["letterhead"] == v1["letterhead"]  # preview keeps showing the letterhead
    blocks = letter["content"]["blocks"]
    table = next(b for b in blocks if b["type"] == "table")
    assert table["rows"][0][2] == "3"
    assert any(b["type"] == "paragraph" and "Added by hand" in b["text"] for b in blocks)
    assert "Contract period" in letter["missing_fields"]  # highlighted marker read back
    # The stored file is hers, byte for byte.
    assert (await client.get(f"/api/v1/demands/{demand['id']}/docx")).content == buf.getvalue()

    # The AI's next edit starts from her version.
    await send(client, demand["id"], "now make the title bold")
    await work(settings)
    context = adapter.requests[0].turns[-1].parts[0].text
    assert context.startswith("CURRENT LETTER (version 2)") and "Added by hand in Word." in context


async def test_upload_edited_letter_rules(client, make_user, admin_password_set, ai, settings):
    first = await letter_ready(client, make_user, admin_password_set, ai, settings)
    original = (await client.get(f"/api/v1/demands/{first['id']}/docx")).content
    other = (await client.post("/api/v1/demands")).json()
    wrong = await client.post(f"/api/v1/demands/{other['id']}/edited-uploads",
                              files={"file": ("x.docx", original, DOCX)})
    assert wrong.status_code == 422 and wrong.json()["type"] == "urn:crm:problem:wrong-letter"
    pdf = await client.post(f"/api/v1/demands/{first['id']}/edited-uploads",
                            files={"file": ("x.pdf", b"%PDF-1.4", "application/pdf")})
    assert pdf.status_code == 415
    broken = await client.post(f"/api/v1/demands/{first['id']}/edited-uploads",
                               files={"file": ("x.docx", b"not a docx", DOCX)})
    assert broken.status_code == 422


async def test_approving_uploaded_letter_keeps_her_body(client, make_user, admin_password_set, ai, settings):
    demand = await letter_ready(client, make_user, admin_password_set, ai, settings)
    doc = Document(io.BytesIO((await client.get(f"/api/v1/demands/{demand['id']}/docx")).content))
    special = doc.add_paragraph("Hand formatting")
    special.runs[0].italic = True
    buf = io.BytesIO()
    doc.save(buf)
    await client.post(f"/api/v1/demands/{demand['id']}/edited-uploads", files={"file": ("e.docx", buf.getvalue(), DOCX)})
    await submit_and_approve(client, demand["id"], 2)
    approved = Document(io.BytesIO((await client.get(f"/api/v1/demands/{demand['id']}/docx")).content))
    last = [p for p in approved.paragraphs if p.text == "Hand formatting"][0]
    assert last.runs[0].italic  # her formatting survived approval


# ---------------- reference letters ----------------


async def test_reference_library_used_unless_chat_has_its_own(client, make_user, admin_password_set, ai, settings):
    adapter = ai(ASK, ASK)
    demand = await setup(client, make_user, admin_password_set)
    ref = Document()
    ref.add_paragraph("REFERENCE FORMAT: Section A — Positions table. Section B — Terms. " * 3)
    buf = io.BytesIO()
    ref.save(buf)
    up = await client.post("/api/v1/admin/reference-letters", data={"label": "Gulf format"},
                           files={"file": ("ref.docx", buf.getvalue(), DOCX)})
    assert up.status_code == 201, up.text

    await send(client, demand["id"], "", [inquiry_file()])
    await work(settings)
    assert "REFERENCE LETTER — Gulf format" in adapter.requests[0].system

    # A reference letter attached in the chat overrides the library.
    await send(client, demand["id"], "use this format", [("mine.docx", buf.getvalue(), DOCX, "custom_reference")])
    await work(settings)
    assert "Gulf format" not in adapter.requests[1].system

    # Deactivated letters are not used; deleting removes it.
    listing = (await client.get("/api/v1/admin/reference-letters")).json()["items"]
    assert listing[0]["characters"] > 50
    assert (await client.patch(f"/api/v1/admin/reference-letters/{listing[0]['id']}", json={"is_active": False})).status_code == 200
    assert (await client.delete(f"/api/v1/admin/reference-letters/{listing[0]['id']}")).status_code == 204
    bad = await client.post("/api/v1/admin/reference-letters", files={"file": ("r.xlsx", b"x", "application/octet-stream")})
    assert bad.status_code == 415


# ---------------- spend cap + usage ----------------


async def test_spend_cap_blocks_new_messages(client, make_user, admin_password_set, ai, settings):
    demand = await letter_ready(client, make_user, admin_password_set, ai, settings)
    usage = (await client.get("/api/v1/admin/usage")).json()
    assert usage["calls"] == 1 and Decimal(usage["total_cost_usd"]) > 0
    assert usage["by_model"][0]["model"] == "claude-sonnet-5-5"

    missing = await client.put("/api/v1/admin/spend-cap", json={"enabled": True})
    assert missing.status_code == 422
    cap = await client.put("/api/v1/admin/spend-cap", json={"enabled": True, "monthly_usd": "0.00"})
    assert cap.status_code == 200 and cap.json()["enabled"] is True

    blocked = await send(client, demand["id"], "one more change")
    assert blocked.status_code == 429
    body = blocked.json()
    assert body["type"] == "urn:crm:problem:spend-cap-reached" and "00:00 UTC" in body["detail"]
    assert int(blocked.headers["retry-after"]) > 0

    off = await client.put("/api/v1/admin/spend-cap", json={"enabled": False})
    assert off.status_code == 200
    ai(LETTER)
    assert (await send(client, demand["id"], "one more change")).status_code == 202


async def test_usage_counts_only_the_utc_month(client, make_user, admin_password_set, ai, settings):
    await letter_ready(client, make_user, admin_password_set, ai, settings)
    async with get_sessionmaker()() as db:
        await db.execute(text("insert into llm_calls (provider, model, outcome, cost_usd, created_at) "
                              "values ('claude', 'claude-sonnet-5-5', 'ok', 5, now() - interval '40 days')"))
        await db.commit()
    now = (await client.get("/api/v1/admin/usage")).json()
    assert Decimal(now["total_cost_usd"]) < 1
    assert (await client.get("/api/v1/admin/usage?month=2026-13")).status_code == 422


async def test_spend_and_usage_need_admin(client, make_user):
    await make_user("ravi")
    await login(client, "ravi")
    assert (await client.get("/api/v1/admin/usage")).status_code == 403
    assert (await client.put("/api/v1/admin/spend-cap", json={"enabled": False})).status_code == 403


# ---------------- full flow with each of the client's sample inquiries ----------------

import pytest  # noqa: E402

SAMPLES = [
    ("inquiry1_mosanada.docx", "jd_list_of_resources.xlsx", "Project: Katara Cultural Village"),
    ("inquiry2_axion.docx", None, "GRP Pipe Fitter"),
    ("inquiry3_mastro.docx", None, "Senior HVAC Technician"),
    ("inquiry4_hvac.docx", "jd_sla_kenya.xlsx", "4000+FAT"),
]
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.mark.parametrize("inquiry,jd,expected_text", SAMPLES)
async def test_full_flow_with_client_samples(client, make_user, admin_password_set, ai, settings,
                                             inquiry, jd, expected_text):
    adapter = ai(ASK, LETTER)
    demand = await setup(client, make_user, admin_password_set)
    files = [inquiry_file(inquiry)]
    if jd:
        files.append((jd, (FIX / jd).read_bytes(), XLSX, "jd"))
    assert (await send(client, demand["id"], "Prepare the demand letter", files)).status_code == 202
    await work(settings)
    sent = "\n".join(p.text or "" for p in adapter.requests[0].turns[0].parts)
    assert expected_text in sent  # the AI received the inquiry's content

    await send(client, demand["id"], "Go ahead with what you have.")
    await work(settings)
    doc = (await client.get(f"/api/v1/demands/{demand['id']}/docx")).content
    edited = Document(io.BytesIO(doc))
    edited.add_paragraph("Checked by the team.")
    buf = io.BytesIO()
    edited.save(buf)
    assert (await client.post(f"/api/v1/demands/{demand['id']}/edited-uploads",
                              files={"file": ("edited.docx", buf.getvalue(), DOCX)})).status_code == 201
    approve = await submit_and_approve(client, demand["id"], 2)
    assert approve.status_code == 200
    final = (await client.get(f"/api/v1/demands/{demand['id']}/docx")).content
    assert "DRAFT" not in header_xml(final) and demand["demand_code"] in header_xml(final)
    assert "Checked by the team." in [p.text for p in Document(io.BytesIO(final)).paragraphs]
    row = (await client.get("/api/v1/demands")).json()["items"][0]
    assert row["demand_code"] == demand["demand_code"] and row["latest_version_no"] == 3
